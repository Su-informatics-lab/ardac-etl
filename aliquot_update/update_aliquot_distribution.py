#!/usr/bin/env python3
"""Create Gen3 aliquot updates from biospecimen distribution manifests."""

from __future__ import annotations

import argparse
import csv
import logging
import sys
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - exercised only in an unconfigured environment
    load_workbook = None


LOGGER = logging.getLogger("aliquot_update")

CATALOG_COLUMNS = {
    "Collection Protocol",
    "PPID",
    "External Subject ID",
    "Visit Name",
    "Barcode",
    "Specimen Type",
    "Anatomic Site",
    "Initial Quantity",
    "Available Quantity",
    "Availability Status",
    "Freeze/Thaw Cycles",
    "Collection Container",
    "Container Display Name",
    "Container Position",
    "Distribution Protocol Short Title",
    "Distribution Date",
}

MANIFEST_COLUMNS = {
    "Collection Protocol Short Title",
    "PPID",
    "External Subject ID",
    "Visit Name",
    "Visit Event Label",
    "Specimen Barcode",
    "Specimen Type",
    "Specimen Request ID",
    "Box",
    "Box Position",
}

ALIQUOT_COLUMNS = {
    "type",
    "id",
    "project_id",
    "submitter_id",
    "aliquot_amount",
    "aliquot_collection_unit",
    "container_type",
    "specimen_type",
    "follow_ups.id",
    "follow_ups.submitter_id",
    "labs.id",
    "labs.submitter_id",
}

LAB_COLUMNS = {
    "type",
    "id",
    "project_id",
    "submitter_id",
    "name_of_institute",
    "PI_email",
    "PI_name",
    "address",
    "center_type",
    "code",
    "contact_email",
    "contact_name",
    "description_of_lab",
    "keyword_name",
    "name",
    "namespace",
    "note",
    "short_name",
    "translational_projects_title",
    "url_of_lab_website",
    "projects.id",
    "projects.code",
}

OUTPUT_SOURCE_COLUMNS = {
    "type",
    "project_id",
    "submitter_id",
    "follow_ups.submitter_id",
    "labs.submitter_id",
    "aliquot_amount",
    "aliquot_collection_unit",
    "container_type",
    "specimen_type",
}

# Extend this table as new Indiana Biobank distribution protocols are introduced.
# Keys contain the institute and PI text from "IB:{institute}:{PI}". Values
# contain the corresponding (name_of_institute, PI_name) in the Gen3 lab dump.
PI_MAPPING: dict[tuple[str, str], tuple[str, str]] = {
    ("Mayo Clinic", "Douglas Simonetto"): ("Mayo Clinic Rochester", "Simonetto"),
}

QC_FILENAMES = {
    "missing_catalog": "missing_catalog_biospecimens.tsv",
    "extraneous_protocol": "extraeneous_collection_protocols.tsv",
    "missing_ardac": "missing_ardac_biospecimens.tsv",
    "redundant": "redundant_manifest_records.tsv",
    "inconsistent_availability": "inconsistent_availability_status.tsv",
}


class InputError(ValueError):
    """Raised when an input file or record cannot be processed safely."""


@dataclass(frozen=True)
class Table:
    columns: tuple[str, ...]
    rows: tuple[dict[str, str], ...]


@dataclass(frozen=True)
class UpdateResult:
    updated_count: int
    qc_counts: dict[str, int]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(path: Path) -> Table:
    """Read a TSV or the active sheet of an XLSX/XLSM workbook."""
    if not path.is_file():
        raise InputError(f"Input file does not exist: {path}")

    suffix = path.suffix.lower()
    if suffix == ".tsv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            if reader.fieldnames is None:
                raise InputError(f"Input file has no header: {path}")
            columns = tuple(name.strip() for name in reader.fieldnames)
            if len(columns) != len(set(columns)):
                raise InputError(f"Input file has duplicate column names: {path}")
            rows = tuple(
                {column: (row.get(column) or "").strip() for column in columns}
                for row in reader
                if any((row.get(column) or "").strip() for column in columns)
            )
            return Table(columns, rows)

    if suffix not in {".xlsx", ".xlsm"}:
        raise InputError(f"Unsupported file type for {path}; expected .tsv, .xlsx, or .xlsm")
    if load_workbook is None:
        raise InputError("Excel input requires openpyxl; install dependencies first")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        try:
            raw_header = next(values)
        except StopIteration as exc:
            raise InputError(f"Input workbook has no header: {path}") from exc
        columns = tuple(_cell_text(value) for value in raw_header)
        while columns and not columns[-1]:
            columns = columns[:-1]
        if not columns or any(not column for column in columns):
            raise InputError(f"Input workbook has an empty column name: {path}")
        if len(columns) != len(set(columns)):
            raise InputError(f"Input workbook has duplicate column names: {path}")
        rows: list[dict[str, str]] = []
        for values_row in values:
            row = {
                column: _cell_text(values_row[index]) if index < len(values_row) else ""
                for index, column in enumerate(columns)
            }
            if any(row.values()):
                rows.append(row)
        return Table(columns, tuple(rows))
    finally:
        workbook.close()


def require_columns(table: Table, required: set[str], label: str, path: Path) -> None:
    missing = sorted(required.difference(table.columns))
    if missing:
        raise InputError(f"{label} {path} is missing required columns: {', '.join(missing)}")


def unique_index(
    rows: Iterable[dict[str, str]], key_column: str, label: str
) -> dict[str, dict[str, str]]:
    result: dict[str, dict[str, str]] = {}
    duplicates: set[str] = set()
    for row in rows:
        key = row[key_column].strip()
        if not key:
            raise InputError(f"{label} contains a record with an empty {key_column}")
        if key in result:
            duplicates.add(key)
        result[key] = row
    if duplicates:
        preview = ", ".join(sorted(duplicates)[:10])
        raise InputError(f"{label} contains duplicate {key_column} values: {preview}")
    return result


def _lab_lookup(labs: Iterable[dict[str, str]]) -> dict[tuple[str, str], str]:
    lookup: dict[tuple[str, str], str] = {}
    for row in labs:
        key = (row["name_of_institute"].strip(), row["PI_name"].strip())
        submitter_id = row["submitter_id"].strip()
        if key in lookup and lookup[key] != submitter_id:
            raise InputError(
                "Lab records do not uniquely map institute/PI pair "
                f"{key!r} to a submitter_id"
            )
        lookup[key] = submitter_id
    return lookup


def _distribution_lab_id(protocol: str, lab_lookup: dict[tuple[str, str], str]) -> str:
    parts = [part.strip() for part in protocol.split(":")]
    if len(parts) != 3 or parts[0] != "IB" or not parts[1] or not parts[2]:
        raise InputError(
            f"Invalid Distribution Protocol Short Title {protocol!r}; "
            "expected IB:{Institute Name}:{PI Name}"
        )
    source_pair = (parts[1], parts[2])
    try:
        lab_pair = PI_MAPPING[source_pair]
    except KeyError as exc:
        raise InputError(
            f"No PI mapping is configured for distribution protocol {protocol!r}"
        ) from exc
    try:
        return lab_lookup[lab_pair]
    except KeyError as exc:
        raise InputError(
            f"Mapped lab {lab_pair!r} is absent from the supplied lab records"
        ) from exc


def _write_single_column(path: Path, barcodes: set[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["barcode"])
        writer.writerows([barcode] for barcode in sorted(barcodes))


def _template_source_name(header: str) -> str:
    return header.removeprefix("*")


def _write_updates(
    path: Path,
    template_columns: Sequence[str],
    rows: Iterable[dict[str, str]],
) -> int:
    count = 0
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=list(template_columns), delimiter="\t", lineterminator="\n"
        )
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {header: row.get(_template_source_name(header), "") for header in template_columns}
            )
            count += 1
    return count


def update_aliquot_distribution(
    *,
    study: str,
    catalog_path: Path,
    manifest_paths: Sequence[Path],
    aliquot_path: Path,
    lab_path: Path,
    output_template_path: Path,
    output_directory: Path,
) -> UpdateResult:
    """Validate inputs, write QC reports, and generate ``updated_aliquot.tsv``."""
    if not study.strip():
        raise InputError("Study collection protocol must not be empty")
    if not manifest_paths:
        raise InputError("At least one distribution manifest is required")

    LOGGER.info("Reading and validating input schemas")
    catalog = read_table(catalog_path)
    require_columns(catalog, CATALOG_COLUMNS, "Catalog", catalog_path)
    aliquots = read_table(aliquot_path)
    require_columns(aliquots, ALIQUOT_COLUMNS, "Aliquot file", aliquot_path)
    labs = read_table(lab_path)
    require_columns(labs, LAB_COLUMNS, "Lab file", lab_path)
    template = read_table(output_template_path)
    template_sources = {_template_source_name(column) for column in template.columns}
    missing_template_columns = sorted(OUTPUT_SOURCE_COLUMNS.difference(template_sources))
    if missing_template_columns:
        raise InputError(
            f"Output template {output_template_path} is missing required columns: "
            + ", ".join(missing_template_columns)
        )

    manifests: list[dict[str, str]] = []
    for manifest_path in manifest_paths:
        manifest = read_table(manifest_path)
        require_columns(manifest, MANIFEST_COLUMNS, "Manifest", manifest_path)
        manifests.extend(manifest.rows)

    catalog_by_barcode = unique_index(catalog.rows, "Barcode", "Catalog")
    aliquot_by_barcode = unique_index(aliquots.rows, "submitter_id", "Aliquot file")
    lab_lookup = _lab_lookup(labs.rows)

    qc: dict[str, set[str]] = {category: set() for category in QC_FILENAMES}
    manifest_barcodes: set[str] = set()
    for row in manifests:
        barcode = row["Specimen Barcode"].strip()
        if not barcode:
            # Manifests can include explicitly marked empty box positions. They
            # are layout metadata, not biospecimen records.
            LOGGER.info("Ignoring a manifest row with no Specimen Barcode")
            continue
        manifest_barcodes.add(barcode)
        if row["Collection Protocol Short Title"].strip() != study:
            qc["extraneous_protocol"].add(barcode)
        catalog_row = catalog_by_barcode.get(barcode)
        if catalog_row is None:
            qc["missing_catalog"].add(barcode)
        else:
            distributed = catalog_row["Availability Status"].strip().casefold() == "distributed"
            has_date = bool(catalog_row["Distribution Date"].strip())
            if not distributed or not has_date:
                qc["inconsistent_availability"].add(barcode)
        aliquot_row = aliquot_by_barcode.get(barcode)
        if aliquot_row is None:
            qc["missing_ardac"].add(barcode)
        elif aliquot_row["labs.submitter_id"].strip() not in {"", "lab_0"}:
            qc["redundant"].add(barcode)

    output_directory.mkdir(parents=True, exist_ok=True)
    for category, filename in QC_FILENAMES.items():
        report_path = output_directory / filename
        _write_single_column(report_path, qc[category])
        LOGGER.info("Generated %s with %d entries", report_path, len(qc[category]))

    rejected = set().union(*qc.values())
    updated_rows: list[dict[str, str]] = []
    for barcode in sorted(manifest_barcodes.difference(rejected)):
        catalog_row = catalog_by_barcode[barcode]
        updated = dict(aliquot_by_barcode[barcode])
        updated["labs.submitter_id"] = _distribution_lab_id(
            catalog_row["Distribution Protocol Short Title"].strip(), lab_lookup
        )
        updated_rows.append(updated)

    output_path = output_directory / "updated_aliquot.tsv"
    updated_count = _write_updates(output_path, template.columns, updated_rows)
    LOGGER.info("Generated %s with %d updated aliquots", output_path, updated_count)
    return UpdateResult(updated_count, {key: len(value) for key, value in qc.items()})


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Update Gen3 aliquot laboratory links from distribution manifests."
    )
    parser.add_argument("--study", required=True, help="Expected study collection protocol")
    parser.add_argument("--catalog", required=True, type=Path, help="Biobank catalog TSV/XLSX")
    parser.add_argument(
        "--manifest",
        required=True,
        type=Path,
        nargs="+",
        help="One or more distribution manifest TSV/XLSX files",
    )
    parser.add_argument("--aliquots", required=True, type=Path, help="Current Gen3 aliquot dump")
    parser.add_argument("--labs", required=True, type=Path, help="Current Gen3 lab dump")
    parser.add_argument(
        "--output-template", required=True, type=Path, help="Gen3 aliquot submission template"
    )
    parser.add_argument("--output-dir", required=True, type=Path, help="Output directory")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    args = build_parser().parse_args(argv)
    try:
        update_aliquot_distribution(
            study=args.study,
            catalog_path=args.catalog,
            manifest_paths=args.manifest,
            aliquot_path=args.aliquots,
            lab_path=args.labs,
            output_template_path=args.output_template,
            output_directory=args.output_dir,
        )
    except (InputError, OSError) as exc:
        LOGGER.error("%s", exc)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
